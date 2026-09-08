"""Indexed artifacts, safe output paths and bounded document previews."""
import csv
import hashlib
import json
from pathlib import Path


def output_directory(repo, task_id, settings):
    base = Path(settings['output_dir']).expanduser().resolve() / task_id if settings.get('output_dir') else repo.directory(task_id).resolve()
    relative = Path(settings.get('subdirectory') or '')
    if relative.is_absolute() or '..' in relative.parts or ':' in str(relative): raise ValueError('子目录必须位于输出目录内')
    directory = (base / relative).resolve()
    if not directory.is_relative_to(base): raise ValueError('输出位置超出运行目录')
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def export_step(service, run, step):
    repo = service.repo
    settings = run.step_settings[step.id]
    source_id = run.workflow.history_id if step.input == 'source' and run.workflow.source == 'history' else run.id
    result = repo.results(source_id)
    input_step = next((s for s in run.workflow.steps if s.id == step.input), None)
    rows = [] if input_step and input_step.kind in ('media','files','export') else [r for r in result['records'] if step.input == 'source' or r.get('step_id') == step.input]
    if not rows: rows = [r for r in result['files'] if step.input == 'source' or r.get('step_id') == step.input]
    rows = [{k:v for k,v in r.items() if k not in ('url','headers','streams','segments','path','external')} for r in rows]
    fields = [key.strip() for key in settings.get('export_fields', '').split(',') if key.strip()]
    if fields: rows = [{key: row.get(key, (row.get('fields') or {}).get(key, '')) for key in fields if key not in ('url','headers','streams','segments','path','external')} for row in rows]
    output = settings['output']
    suffix = 'xlsx' if output == 'excel' else output
    name = settings.get('filename') or 'results'
    if Path(name).name != name or ':' in name or name in ('.','..'): raise ValueError('文件名称无效')
    directory = output_directory(repo, run.id, settings) / step.id
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f'{name}.{suffix}'
    if path.exists(): raise ValueError('导出文件已存在，请更换名称')
    if output == 'json': path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
    elif output == 'jsonl': path.write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows), encoding='utf-8')
    else:
        columns = list(dict.fromkeys(k for row in rows for k in row)) or ['title']
        values = [[json.dumps(r.get(k,''),ensure_ascii=False) if isinstance(r.get(k), (dict,list)) else str(r.get(k,'')) for k in columns] for r in rows]
        if output == 'csv':
            with path.open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.writer(stream);writer.writerow(columns)
                writer.writerows([["'"+v if v.startswith(('=','+','-','@')) else v for v in row] for row in values])
        else:
            from openpyxl import Workbook
            book=Workbook();sheet=book.active;sheet.append(columns)
            for row in values: sheet.append(row)
            for row in sheet:
                for cell in row: cell.data_type='s'
            book.save(path);book.close()
    repo.file(run.id,dict(id='export-'+step.id, step_id=step.id, title=path.name, kind='export',state='succeeded',path=str(path.resolve()),external=True,size=path.stat().st_size,error=''))


def artifact_rows(repo, task_id):
    rows=[]
    for f in repo.results(task_id)['files']:
        rows.append({k:v for k,v in {**f,'task_id':task_id,'extension':Path(f.get('path') or f.get('title','')).suffix.lstrip('.')}.items() if k not in ('url','headers','streams','segments','path','external')})
    for p in repo.directory(task_id).glob('results.*'):
        if p.is_file(): rows.append(dict(id='legacy-'+hashlib.sha256(p.name.encode()).hexdigest()[:20],task_id=task_id,step_id='legacy',title=p.name,kind='export',state='succeeded',size=p.stat().st_size))
    return rows


def artifact_path(repo, task_id, artifact_id):
    repo.task(task_id)
    for f in repo.results(task_id)['files']:
        if f['id'] == artifact_id and f['state']=='succeeded':
            owner=f.get('source_task_id') or task_id
            root=repo.directory(owner).resolve()
            path=Path(f['path']) if f.get('external') else root/f['path']
            path=path.resolve()
            if not f.get('external') and not path.is_relative_to(root): break
            if path.is_file(): return path
    for p in repo.directory(task_id).glob('results.*'):
        if 'legacy-'+hashlib.sha256(p.name.encode()).hexdigest()[:20] == artifact_id and p.is_file(): return p
    raise ValueError('文件不存在或尚未下载完成')


def document_preview(path, offset=0, limit=100):
    suffix=path.suffix.lower()
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            from itertools import islice
            rows=[[str(v)[:10000] if v is not None else '' for v in row[:100]] for row in islice(book.active.values,offset,offset+limit+1)]
            return dict(type='table',rows=rows[:limit],more=len(rows)>limit,offset=offset)
        finally:book.close()
    if suffix=='.csv':
        with path.open(encoding='utf-8-sig',errors='replace',newline='') as stream:
            from itertools import islice
            rows=[[cell[:10000] for cell in row[:100]] for row in islice(csv.reader(stream),offset,offset+limit+1)]
        return dict(type='table',rows=[r[:100] for r in rows[:limit]],more=len(rows)>limit,offset=offset)
    if suffix in ('.txt','.json','.jsonl','.log','.md'):
        with path.open('rb') as stream:
            stream.seek(offset);data=stream.read(65537)
        return dict(type='text',text=data[:65536].decode('utf-8',errors='replace'),more=len(data)>65536,offset=offset,next_offset=offset+65536)
    return dict(type='unsupported',more=False)



def file_response(path, download=False, byte_range=None):
    """Single-range streaming for the pinned Starlette version's media viewers."""
    import re
    from fastapi.responses import FileResponse, StreamingResponse, Response
    inline=path.suffix.lower() in ('.pdf','.mp4','.webm','.mp3','.m4a','.ogg','.wav','.jpg','.jpeg','.png','.gif','.webp','.avif')
    response=FileResponse(path,filename=path.name,content_disposition_type='inline' if inline and not download else 'attachment',headers={'X-Content-Type-Options':'nosniff','Accept-Ranges':'bytes'})
    if not byte_range:return response
    size=path.stat().st_size
    try:
        match=re.fullmatch(r'bytes=(\d*)-(\d*)',byte_range.strip())
        if not match or not size or not any(match.groups()):raise ValueError()
        first,last=match.groups()
        if first:
            start=int(first);end=min(int(last) if last else size-1,size-1)
        else:
            suffix=int(last)
            if suffix<=0:raise ValueError()
            start=max(0,size-suffix);end=size-1
        if start>=size or end<start:raise ValueError()
    except ValueError:return Response(status_code=416,headers={'Content-Range':f'bytes */{size}','Accept-Ranges':'bytes','X-Content-Type-Options':'nosniff'})
    def chunks():
        with path.open('rb') as stream:
            stream.seek(start);remaining=end-start+1
            while remaining:
                data=stream.read(min(65536,remaining))
                if not data:break
                remaining-=len(data)
                yield data
    return StreamingResponse(chunks(),status_code=206,media_type=response.media_type,
        headers={**response.headers,'Content-Length':str(end-start+1),'Content-Range':f'bytes {start}-{end}/{size}'})
